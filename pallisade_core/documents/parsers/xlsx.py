from __future__ import annotations

import io
from dataclasses import dataclass

from openpyxl import load_workbook


@dataclass(frozen=True)
class SheetText:
    sheet_name: str
    text: str
    row_count: int


def extract_sheets(data: bytes) -> list[SheetText]:
    """Return a list of (sheet_name, flattened_text, row_count) per sheet.

    Each sheet's text starts with a header line ``[Sheet: <name>, rows N-M]``
    followed by one line per row formatted as ``A1: ... | B1: ... | ...``.
    Empty rows and empty cells are skipped.
    """
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    out: list[SheetText] = []
    try:
        for ws in wb.worksheets:
            rows = []
            min_row = None
            max_row = 0
            for r_idx, row in enumerate(ws.iter_rows(values_only=False), start=1):
                cells = []
                for cell in row:
                    if cell.value is None or cell.value == "":
                        continue
                    cells.append(f"{cell.coordinate}: {cell.value}")
                if cells:
                    rows.append(" | ".join(cells))
                    if min_row is None:
                        min_row = r_idx
                    max_row = r_idx
            if not rows:
                continue
            header = f"[Sheet: {ws.title}, rows {min_row}-{max_row}]"
            body = "\n".join(rows)
            out.append(
                SheetText(
                    sheet_name=ws.title,
                    text=f"{header}\n{body}",
                    row_count=len(rows),
                )
            )
    finally:
        wb.close()
    return out
